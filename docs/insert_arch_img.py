#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""直接在已格式化的 docx 中替换 ASCII art 为架构图 PNG"""
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document("/workspace/docs/毕业设计论文.docx")

# 找到第一个含 ASCII art 的段落，替换为图片
replaced = False
paras_to_clear = []

for i, p in enumerate(doc.paragraphs):
    text = p.text
    if '┌' in text or '└' in text or '│' in text or '▼' in text or '├' in text:
        if not replaced:
            # 清空该段落内容并插入图片
            for run in p.runs:
                run.text = ""
            # 添加新 run 并插入图片
            run = p.add_run()
            run.add_picture("/workspace/docs/arch_diagram.png", width=Inches(5.5))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            replaced = True
        else:
            # 清空其余 ASCII 行
            paras_to_clear.append(p)

for p in paras_to_clear:
    for run in p.runs:
        run.text = ""

doc.save("/workspace/docs/毕业设计论文.docx")

# 核验
doc2 = Document("/workspace/docs/毕业设计论文.docx")
img_count = sum(1 for rel in doc2.part.rels.values() if 'image' in rel.reltype)
ascii_count = sum(1 for p in doc2.paragraphs if any(c in p.text for c in '┌└│▼├'))
print(f"图片数量: {img_count}")
print(f"ASCII art 残留: {ascii_count}")
print(f"架构图替换: {'成功' if replaced and img_count > 0 else '失败'}")
