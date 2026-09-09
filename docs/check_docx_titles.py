#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""核验 Word 中所有标题的样式层级和内容"""
from docx import Document
from docx.oxml.ns import qn
import re

doc = Document("/workspace/docs/毕业设计论文.docx")

print("=" * 80)
print("Word 标题层级核验")
print("=" * 80)
print(f"{'行号':<6}{'样式':<15}{'文本前40字符'}")
print("-" * 80)

for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if not text:
        continue
    style = p.style.name if p.style else "无"

    # 显示所有标题和疑似标题的段落
    if (style.startswith("Heading") or
        re.match(r'^第\d+章', text) or
        re.match(r'^\d+\.\d+', text) or
        text in ("摘要", "Abstract", "目录", "结论", "参考文献", "致谢", "附录", "诚信承诺书") or
        text.startswith("关键词") or text.startswith("Keywords") or
        re.match(r'^(表|图)\s', text) or
        re.match(r'^附录', text)):
        # 获取第一个run的字体信息
        info = ""
        if p.runs:
            run = p.runs[0]
            size = run.font.size.pt if run.font.size else "默认"
            bold = run.font.bold
            rpr = run._element.find(qn('w:rPr'))
            cn_font = ""
            if rpr is not None:
                rfonts = rpr.find(qn('w:rFonts'))
                if rfonts is not None:
                    cn_font = rfonts.get(qn('w:eastAsia'), "")
            info = f" [{cn_font} {size}pt {'加粗' if bold else '常规'}]"
        print(f"{i:<6}{style:<15}{text[:40]}{info}")

print("-" * 80)
print(f"段落总数: {len(doc.paragraphs)}")
