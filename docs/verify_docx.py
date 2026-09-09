#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""核验 Word 字体字号设置是否符合规范"""
from docx import Document
from docx.oxml.ns import qn
import re

doc = Document("/workspace/docs/毕业设计论文.docx")

# 字号对照表
size_map = {16: "三号", 18: "小二号", 15: "小三号", 14: "四号", 12: "小四号", 10.5: "五号"}

print("=" * 70)
print("Word 字体字号核验报告")
print("=" * 70)
print(f"段落总数: {len(doc.paragraphs)}")
print(f"表格总数: {len(doc.tables)}")
print("-" * 70)

# 抽样核验各层级
samples = {"三号标题": 0, "小三标题": 0, "四号标题": 0, "小四正文": 0, "五号表内": 0, "摘要": 0, "Abstract": 0, "表标题": 0}

for p in doc.paragraphs:
    text = p.text.strip()
    if not p.runs:
        continue
    run = p.runs[0]
    size = run.font.size.pt if run.font.size else "默认"
    bold = run.font.bold
    cn_font = ""
    rpr = run._element.find(qn('w:rPr'))
    if rpr is not None:
        rfonts = rpr.find(qn('w:rFonts'))
        if rfonts is not None:
            cn_font = rfonts.get(qn('w:eastAsia'), "")

    # 核验关键段落
    if re.match(r'^第\d+章', text) and samples["三号标题"] == 0:
        print(f"[章标题] {text[:30]}")
        print(f"  字号={size}pt({size_map.get(size, '?')})  字体={cn_font}  加粗={bold}")
        samples["三号标题"] = 1
    elif re.match(r'^\d+\.\d+\s', text) and samples["小三标题"] == 0:
        print(f"[二级标题] {text[:30]}")
        print(f"  字号={size}pt({size_map.get(size, '?')})  字体={cn_font}  加粗={bold}")
        samples["小三标题"] = 1
    elif re.match(r'^\d+\.\d+\.\d+\s', text) and samples["四号标题"] == 0:
        print(f"[三级标题] {text[:30]}")
        print(f"  字号={size}pt({size_map.get(size, '?')})  字体={cn_font}  加粗={bold}")
        samples["四号标题"] = 1
    elif text == "摘要" and samples["摘要"] == 0:
        print(f"[中文摘要标题] {text}")
        print(f"  字号={size}pt({size_map.get(size, '?')})  字体={cn_font}  加粗={bold}")
        samples["摘要"] = 1
    elif text == "Abstract" and samples["Abstract"] == 0:
        print(f"[英文摘要标题] {text}")
        print(f"  字号={size}pt({size_map.get(size, '?')})  字体={cn_font}  加粗={bold}")
        samples["Abstract"] = 1
    elif re.match(r'^表\s*\S', text) and samples["表标题"] == 0:
        print(f"[表标题] {text[:30]}")
        print(f"  字号={size}pt({size_map.get(size, '?')})  字体={cn_font}  加粗={bold}")
        samples["表标题"] = 1

# 核验表格内文字
if doc.tables:
    t = doc.tables[0]
    cell = t.rows[0].cells[0]
    for p in cell.paragraphs:
        if p.runs:
            run = p.runs[0]
            size = run.font.size.pt if run.font.size else "默认"
            rpr = run._element.find(qn('w:rPr'))
            cn_font = ""
            if rpr is not None:
                rfonts = rpr.find(qn('w:rFonts'))
                if rfonts is not None:
                    cn_font = rfonts.get(qn('w:eastAsia'), "")
            print(f"[表内文字] {p.text[:20]}")
            print(f"  字号={size}pt({size_map.get(size, '?')})  字体={cn_font}")
            break

# 核验普通正文
for p in doc.paragraphs:
    text = p.text.strip()
    if len(text) > 30 and p.runs and not re.match(r'^[#\d]', text) and not re.match(r'^(第|表|图|摘要|Abstract|关键词|Keywords)', text):
        run = p.runs[0]
        size = run.font.size.pt if run.font.size else "默认"
        rpr = run._element.find(qn('w:rPr'))
        cn_font = ""
        if rpr is not None:
            rfonts = rpr.find(qn('w:rFonts'))
            if rfonts is not None:
                cn_font = rfonts.get(qn('w:eastAsia'), "")
        indent = p.paragraph_format.first_line_indent
        indent_pt = indent.pt if indent else "无"
        print(f"[正文段落] {text[:30]}...")
        print(f"  字号={size}pt({size_map.get(size, '?')})  字体={cn_font}  首行缩进={indent_pt}pt")
        break

print("-" * 70)
print("核验完成")
