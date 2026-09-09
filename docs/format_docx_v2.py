#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
修复版 Word 格式化脚本：
1. 重新从 markdown 生成 docx（确保样式干净）
2. 修复标题层级（摘要/Abstract/诚信承诺书/目录等设为 Heading 1）
3. 目录内容行设为普通文本
4. 嵌入架构图 PNG 替换 ASCII art
5. 按规范设置所有字体字号
"""
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_run_font(run, cn_font="宋体", en_font="Times New Roman", size_pt=12, bold=False):
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.name = en_font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    rfonts.set(qn('w:ascii'), en_font)
    rfonts.set(qn('w:hAnsi'), en_font)
    rfonts.set(qn('w:eastAsia'), cn_font)


def set_para(p, alignment=None, first_line_indent=None, line_spacing=1.5,
             space_before=None, space_after=None):
    pf = p.paragraph_format
    if alignment is not None:
        p.alignment = alignment
    if first_line_indent is not None:
        pf.first_line_indent = Pt(first_line_indent)
    pf.line_spacing = line_spacing
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    if space_before is not None:
        pf.space_before = Pt(space_before)
    if space_after is not None:
        pf.space_after = Pt(space_after)


def set_style(p, style_name):
    """设置段落样式，使用 pPr 中的 pStyle 直接设置以确保生效"""
    try:
        # 先尝试用对象 API 设置
        p.style = p.part.document.styles[style_name]
    except (KeyError, Exception):
        pass
    # 双重保险：直接操作 XML 设置 pStyle
    pPr = p._element.get_or_add_pPr()
    # 移除已有的 pStyle
    for old in pPr.findall(qn('w:pStyle')):
        pPr.remove(old)
    pStyle = OxmlElement('w:pStyle')
    # 样式 ID 映射
    style_id_map = {
        "Heading 1": "Heading1",
        "Heading 2": "Heading2",
        "Heading 3": "Heading3",
        "Heading 4": "Heading4",
        "Normal": "Normal",
    }
    pStyle.set(qn('w:val'), style_id_map.get(style_name, style_name))
    pPr.insert(0, pStyle)


def main():
    # 先从 markdown 重新生成干净的 docx
    import subprocess
    subprocess.run([
        "pandoc", "毕业设计论文.md", "-o", "毕业设计论文.docx",
        "--from", "markdown", "--to", "docx"
    ], check=True, cwd="/workspace/docs")

    doc = Document("/workspace/docs/毕业设计论文.docx")

    # 字号常量
    SAN_HAO = 16
    XIAO_ER = 18
    XIAO_SAN = 15
    SI_HAO = 14
    XIAO_SI = 12
    WU_HAO = 10.5

    # 需要设为 Heading 1 的独立章节标题
    h1_titles = {"摘要", "Abstract", "诚信承诺书", "目录", "结论",
                 "参考文献", "致谢", "附录"}

    # 目录内容特征：必须含 "…" 才算目录行
    def is_toc_line(text):
        return '…' in text

    # 是否为 ASCII art 架构图代码块
    in_code_block = False
    arch_diagram_inserted = False

    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        style = p.style.name if p.style else ""

        # 检测代码块开始/结束
        if text.startswith("```"):
            if not in_code_block:
                in_code_block = True
                # 如果是架构图代码块（下一个内容含 ┌ 或 └）
                # 标记为待替换
                p.text = ""  # 清空 ``` 行
                # 检查是否是架构图区域
            else:
                in_code_block = False
                p.text = ""  # 清空 ``` 行
            continue

        # 代码块内的内容 - 如果是架构图则替换为图片
        if in_code_block and ('┌' in text or '└' in text or '│' in text or '▼' in text):
            if not arch_diagram_inserted:
                # 插入架构图图片
                run = p.add_run()
                run.add_picture("/workspace/docs/arch_diagram.png", width=Inches(5.5))
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                arch_diagram_inserted = True
            else:
                p.text = ""  # 清空其余 ASCII 行
            continue

        # 代码块内其他内容（protobuf、目录树等）- 保留为代码格式
        if in_code_block:
            set_para(p, alignment=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.0)
            for run in p.runs:
                set_run_font(run, cn_font="宋体", en_font="Courier New",
                             size_pt=WU_HAO, bold=False)
            continue

        if not text:
            continue

        # === 标题层级修复 ===

        # 论文主标题
        if i == 0 and 'AI驱动' in text:
            set_style(p, "Heading 1")
            set_para(p, alignment=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.5,
                     space_before=12, space_after=12)
            for run in p.runs:
                set_run_font(run, cn_font="黑体", en_font="Times New Roman",
                             size_pt=SAN_HAO, bold=True)

        # 独立章节标题 → Heading 1
        elif text in h1_titles:
            set_style(p, "Heading 1")
            set_para(p, alignment=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.5,
                     space_before=12, space_after=6)
            cn = "黑体" if text == "目录" else "宋体"
            en = "Times New Roman" if text == "Abstract" else "Times New Roman"
            size = XIAO_ER if text == "Abstract" else SAN_HAO
            for run in p.runs:
                set_run_font(run, cn_font=cn, en_font=en, size_pt=size, bold=True)

        # 附录子标题
        elif text.startswith("附录A") or text.startswith("附录 A"):
            set_style(p, "Heading 2")
            set_para(p, alignment=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.5,
                     space_before=8, space_after=4)
            for run in p.runs:
                set_run_font(run, cn_font="宋体", en_font="Times New Roman",
                             size_pt=XIAO_SAN, bold=True)

        # 目录内容行 → 普通文本（非标题）
        elif is_toc_line(text):
            set_style(p, "Normal")
            set_para(p, alignment=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.5,
                     first_line_indent=0)
            for run in p.runs:
                set_run_font(run, cn_font="宋体", en_font="Times New Roman",
                             size_pt=XIAO_SI, bold=False)

        # 第X章 → Heading 2（在正文中作为一级章节标题）
        elif re.match(r'^第\d+章', text):
            set_style(p, "Heading 2")
            set_para(p, alignment=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.5,
                     space_before=12, space_after=6)
            for run in p.runs:
                set_run_font(run, cn_font="宋体", en_font="Times New Roman",
                             size_pt=SAN_HAO, bold=True)

        # 1.1 级 → Heading 3
        elif re.match(r'^\d+\.\d+\s', text):
            set_style(p, "Heading 3")
            set_para(p, alignment=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.5,
                     space_before=8, space_after=4)
            for run in p.runs:
                set_run_font(run, cn_font="宋体", en_font="Times New Roman",
                             size_pt=XIAO_SAN, bold=True)

        # 1.1.1 级 → Heading 4
        elif re.match(r'^\d+\.\d+\.\d+\s', text):
            set_style(p, "Heading 4")
            set_para(p, alignment=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.5,
                     first_line_indent=28, space_before=6, space_after=3)
            for run in p.runs:
                set_run_font(run, cn_font="宋体", en_font="Times New Roman",
                             size_pt=SI_HAO, bold=True)

        # 关键词行（顶格加粗）
        elif text.startswith("关键词") or text.startswith("Keywords"):
            set_style(p, "Normal")
            set_para(p, alignment=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.5,
                     space_before=4)
            for run in p.runs:
                set_run_font(run, cn_font="宋体", en_font="Times New Roman",
                             size_pt=XIAO_SI, bold=True)

        # 表/图标题行
        elif re.match(r'^(表|图)\s*\S', text):
            set_style(p, "Normal")
            set_para(p, alignment=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.5,
                     space_before=4, space_after=2)
            for run in p.runs:
                set_run_font(run, cn_font="黑体", en_font="Times New Roman",
                             size_pt=WU_HAO, bold=True)

        # 普通正文
        else:
            set_style(p, "Normal")
            set_para(p, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.5,
                     first_line_indent=24)
            for run in p.runs:
                set_run_font(run, cn_font="宋体", en_font="Times New Roman",
                             size_pt=XIAO_SI, bold=False)

    # 处理表格内文字
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    set_para(p, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                             line_spacing=1.0, first_line_indent=0)
                    for run in p.runs:
                        set_run_font(run, cn_font="宋体", en_font="Times New Roman",
                                     size_pt=WU_HAO, bold=False)

    # 设置 Normal 默认字体
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(XIAO_SI)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    rfonts.set(qn('w:eastAsia'), '宋体')

    out = "/workspace/docs/毕业设计论文.docx"
    doc.save(out)
    print(f"已保存修复版 Word: {out}")
    print("修复内容: 标题层级统一 + 架构图替换为 PNG + 字体字号规范")


if __name__ == "__main__":
    main()
